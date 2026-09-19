import sys, zipfile, json, hashlib
from pathlib import Path
import openpyxl
root=Path(sys.argv[1]); original=root/"issue-001/orders.xlsx"; demo=root/"issue-002/orders.xlsx"
with zipfile.ZipFile(original) as a, zipfile.ZipFile(demo) as b:
    assert a.namelist()==b.namelist()
    changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
    assert changed==["xl/worksheets/sheet1.xml","xl/worksheets/sheet2.xml"], changed
a=openpyxl.load_workbook(original,data_only=False); b=openpyxl.load_workbook(demo,data_only=False)
deltas=[]
for name in a.sheetnames:
    for row in a[name]:
        for cell in row:
            other=b[name][cell.coordinate]
            assert cell.data_type==other.data_type
            assert cell.number_format==other.number_format
            assert cell._style==other._style
            if cell.value != other.value: deltas.append([name,cell.coordinate,cell.value,other.value])
assert deltas==[["Orders","C2",2,3]], deltas
a.close();b.close()
print(json.dumps({"changedZipMembers":changed,"formulaViewCellDeltas":deltas,"typesFormatsStylesPreserved":True},indent=2))
