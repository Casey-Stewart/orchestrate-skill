import sys
import re
from pathlib import Path
from zipfile import ZipFile
import openpyxl

original, broken = map(Path, sys.argv[1:])
formula = openpyxl.load_workbook(original, data_only=False)
values = openpyxl.load_workbook(original, data_only=True)
assert formula.sheetnames == ["Orders", "Summary", "Types"]
assert formula["Orders"]["A2"].value == "0001"
assert formula["Orders"]["A2"].data_type == "s"
assert formula["Summary"]["B2"].value == "=SUM(Orders!F2:F5)"
assert values["Orders"]["F2"].value == 22.5
assert values["Summary"]["B2"].value == 23.5
assert values["Summary"]["B3"].value == 4
assert isinstance(values["Types"]["B2"].value, bool)
print("Independent openpyxl formula/data-only inspection: PASS; total=23.50 count=4")
if broken.exists():
    raise SystemExit("Refusing to overwrite a scratch corruption fixture")
with ZipFile(original) as source, ZipFile(broken, "w") as destination:
    for info in source.infolist():
        data = source.read(info.filename)
        if info.filename == "xl/worksheets/sheet1.xml":
            data, count = re.subn(rb'(<c r="C2"[^>]*><v>)2(</v></c>)', rb'\g<1>3\g<2>', data)
            assert count == 1
        destination.writestr(info, data)
print("Created disposable Orders!C2 corruption:", broken)
